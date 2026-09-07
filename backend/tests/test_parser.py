"""
Unit tests for Document Parsing Service (Phase 1).
"""

import io
import pytest
import fitz  # PyMuPDF
import docx
from backend.app.services.parser import (
    parse_document,
    parse_pdf,
    parse_docx,
    EmptyDocumentError,
    UnsupportedFileTypeError,
)


def create_sample_pdf(text: str = "This is a sample agreement.") -> bytes:
    """Helper to generate an in-memory PDF for testing."""
    doc = fitz.open()
    page = doc.new_page()
    if text:
        page.insert_text((50, 72), text)
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes


def create_sample_docx(paragraphs: list = None) -> bytes:
    """Helper to generate an in-memory DOCX for testing."""
    doc = docx.Document()
    if paragraphs:
        for p in paragraphs:
            doc.add_paragraph(p)
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_parse_valid_pdf():
    """Verify parsing a valid PDF extracts correct text and metadata."""
    sample_text = "SECTION 1. CONFIDENTIALITY\nThe parties agree to maintain secret all proprietary materials."
    pdf_bytes = create_sample_pdf(sample_text)
    
    result = parse_document(pdf_bytes, "nda_contract.pdf")
    assert result.file_type == "pdf"
    assert result.filename == "nda_contract.pdf"
    assert "CONFIDENTIALITY" in result.text
    assert result.total_characters > 0
    assert len(result.paragraphs) >= 1


def test_parse_empty_pdf_raises_error():
    """Verify empty/blank PDF triggers EmptyDocumentError with helpful message."""
    pdf_bytes = create_sample_pdf("")  # No text inserted
    with pytest.raises(EmptyDocumentError) as exc_info:
        parse_document(pdf_bytes, "scanned_empty.pdf")
    assert "no extractable text" in str(exc_info.value).lower()


def test_parse_valid_docx():
    """Verify parsing a valid DOCX extracts paragraphs and structure."""
    paragraphs = [
        "FREELANCE SERVICE AGREEMENT",
        "1. Scope of Work: Contractor shall deliver web application designs.",
        "2. Compensation: Client agrees to pay $5,000 upon final milestone.",
    ]
    docx_bytes = create_sample_docx(paragraphs)
    
    result = parse_document(docx_bytes, "freelance_agreement.docx")
    assert result.file_type == "docx"
    assert result.filename == "freelance_agreement.docx"
    assert "FREELANCE SERVICE AGREEMENT" in result.text
    assert len(result.paragraphs) == 3


def test_parse_empty_docx_raises_error():
    """Verify blank DOCX triggers EmptyDocumentError."""
    docx_bytes = create_sample_docx([])  # No paragraphs
    with pytest.raises(EmptyDocumentError) as exc_info:
        parse_document(docx_bytes, "blank_contract.docx")
    assert "no extractable text" in str(exc_info.value).lower()


def test_parse_unsupported_file_type():
    """Verify unsupported file extensions are rejected with clear error."""
    with pytest.raises(UnsupportedFileTypeError) as exc_info:
        parse_document(b"fake text content", "contract.txt")
    assert "unsupported file format" in str(exc_info.value).lower()


def test_parse_legacy_doc_file_type():
    """Verify legacy .doc extension provides informative conversion message."""
    with pytest.raises(UnsupportedFileTypeError) as exc_info:
        parse_document(b"legacy word binary", "old_contract.doc")
    assert "convert your file to .docx or .pdf" in str(exc_info.value).lower()
